import subprocess
import time
import uuid
from django.db import models
from django.core.files.storage import FileSystemStorage
from django.conf import settings
from zipfile import ZipFile
import os
from pdf2image import convert_from_path
from docx import Document
from PIL import Image

# Define file storage locations
class OverwriteStorage(FileSystemStorage):
    def get_available_name(self, name, max_length=None):
        file_path = os.path.join(self.location, name)

        if self.exists(name):
            print(f"Attempting to remove: {file_path}")
            attempts = 3  # Nombre de tentatives pour supprimer le fichier
            for attempt in range(attempts):
                try:
                    os.remove(file_path)
                    print(f"File removed: {file_path}")
                    break
                except PermissionError:
                    print(f"Attempt {attempt + 1}/{attempts} failed: {file_path} is locked.")
                    time.sleep(1)  # Attendez 1 seconde avant de réessayer
            else:
                raise PermissionError(f"Cannot remove locked file: {file_path}")
        return name

class Conversion(models.Model):
    input_file = models.FileField(upload_to='uploads/', storage=OverwriteStorage())
    output_file = models.FileField(upload_to='converted/', blank=True, null=True, storage=OverwriteStorage())
    source_format = models.CharField(max_length=10)
    target_format = models.CharField(max_length=10)
    converted = models.BooleanField(default=False)
    error_message = models.CharField(max_length=255, blank=True, null=True)
    token = models.UUIDField(default=uuid.uuid4, editable=False, unique=True)  # Token unique pour chaque conversion

    def __str__(self):
        return f"{self.source_format} to {self.target_format}"

    def convert_file(self):
        print(f"Starting conversion: {self.source_format} to {self.target_format}")

        # Formats supportés
        supported_formats = [
            "png", "jpeg", "jpg", "bmp", "tiff", "gif", "webp", "ico", 
            "docx", "odt", "txt", "xlsx", "pptx", "pdf"
        ]
        
        
        try:
            if self.source_format == self.target_format:
                self.converted = False
                self.error_message = "Source and target formats are the same. No conversion needed."
                self.save()
                return

            if self.source_format.lower() in ["jpg", "jpeg"] and self.target_format.lower() in ["jpg", "jpeg"]:
                self._convert_jpg_to_jpeg()
                return
            
            # Normaliser JPG en JPEG
            if self.source_format.lower() == "jpg":
                self.source_format = "jpeg"
            if self.target_format.lower() == "jpg":
                self.target_format = "jpeg"
                
            if self.source_format == "pdf":
                if self.target_format in supported_formats:
                    self._convert_pdf_to_images()
                else:
                    self._unsupported_format()
                    return
            elif self.target_format == "pdf":
                if self.source_format in ["docx", "odt", "txt", "xlsx", "pptx"]:
                    self._convert_documents_to_pdf()
                elif self.source_format in supported_formats:
                    self._convert_images_to_pdf()
                else:
                    self._unsupported_format()
                    return
            elif self.source_format in ["docx", "odt", "xlsx", "txt"] and self.target_format in ["docx", "odt", "xlsx", "txt"]:
                self._convert_documents_to_document()
            elif self.source_format in supported_formats and self.target_format in supported_formats:
                if self.source_format not in ["docx", "odt", "xlsx", "txt"] and self.target_format not in ["docx", "odt", "xlsx", "txt"]:
                    self._convert_image()
                else:
                    self._unsupported_format()
                    return
            else:
                self._unsupported_format()
                return
            print(f"Conversion {self.source_format} to {self.target_format} completed")
        except Exception as e:
            self.converted = False
            self.error_message = f"General conversion error: {e}"
            self.save()

    def _unsupported_format(self):
        self.converted = False
        self.error_message = f"Unsupported conversion: {self.source_format} to {self.target_format}"
        self.save()

    def _convert_jpg_to_jpeg(self):
        try:
            file_name = os.path.basename(self.input_file.name)
            file_root, _ = os.path.splitext(file_name)
            target_format = self.target_format.lower()
            target_extension = "jpeg" if target_format == "jpeg" else "jpg"
            target_path = os.path.join(settings.MEDIA_ROOT, f"converted/{file_root}.{target_extension}")

            img = Image.open(self.input_file.path).convert("RGB")
            img.save(target_path, "JPEG")  # Toujours utiliser JPEG pour les conversions JPG/JPEG
            self.output_file.name = target_path
            self.converted = True
            self.error_message = None
            self.save()
        except Exception as e:
            self.converted = False
            self.error_message = f"Error during JPG <-> JPEG conversion: {e}"
            self.save()

    def _convert_pdf_to_images(self):
        try:
            poppler_path = None
            if os.name == 'nt':  # Windows
                poppler_path = r"C:/Program Files/poppler-24.08.0/Library/bin"

            temp_dir = os.path.join(settings.MEDIA_ROOT, "temp")
            os.makedirs(temp_dir, exist_ok=True)

            if self.target_format.lower() == "jpg":
                fmt = "jpeg"
            else:
                fmt = self.target_format.lower()
    
            images = convert_from_path(
                self.input_file.path,
                output_folder=temp_dir,
                fmt=fmt,
                poppler_path=poppler_path
            )

            file_name = os.path.basename(self.input_file.name)
            file_root, _ = os.path.splitext(file_name)
            zip_path = os.path.join(settings.MEDIA_ROOT, f"converted/{file_root}.zip")

            with ZipFile(zip_path, 'w') as zipf:
                for i, image in enumerate(images):
                    image_path = os.path.join(temp_dir, f"page_{i + 1}.{self.target_format}")
                    try:
                        image.save(image_path, self.target_format.upper())
                        zipf.write(image_path, os.path.basename(image_path))
                    except Exception as e:
                        self.error_message = f"Error saving image {i + 1}: {e}"
                        self.save()
                        raise

            for file in os.listdir(temp_dir):
                file_path = os.path.join(temp_dir, file)
                if os.path.isfile(file_path):
                    os.remove(file_path)

            self.output_file.name = zip_path
            self.converted = True
            self.save()
        except Exception as e:
            self.converted = False
            self.error_message = f"Error during pdf to {self.target_format} conversion: {e}"
            self.save()

    def _convert_images_to_pdf(self):
        try:
            file_name = os.path.basename(self.input_file.name)
            file_root, _ = os.path.splitext(file_name)
            pdf_path = os.path.join(settings.MEDIA_ROOT, f"converted/{file_root}.pdf")

            img = Image.open(self.input_file.path)
            img.convert("RGB").save(pdf_path, "pdf", resolution=100.0)

            self.output_file.name = pdf_path
            self.converted = True
            self.save()
        except Exception as e:
            self.converted = False
            self.error_message = f"Error during {self.source_format} to pdf conversion: {e}"
            self.save()

    def _convert_image(self):
        try:
            file_name = os.path.basename(self.input_file.name)
            file_root, _ = os.path.splitext(file_name)
            target_path = os.path.join(settings.MEDIA_ROOT, f"converted/{file_root}.{self.target_format}")

            if self.target_format.lower() == "jpg":
                target_format_upper = "JPEG"
            else:
                target_format_upper = self.target_format.upper()

            img = Image.open(self.input_file.path).convert("RGB")
            img.save(target_path, target_format_upper)
            self.output_file.name = target_path
            self.converted = True
            self.save()
        except Exception as e:
            self.converted = False
            self.error_message = f"Error during image conversion: {e}"
            self.save()
            
    def _convert_documents_to_pdf(self):
        try:
            file_name = os.path.basename(self.input_file.name)
            file_root, _ = os.path.splitext(file_name)
            pdf_path = os.path.join(settings.MEDIA_ROOT, f"converted/{file_root}.pdf")

            if self.source_format == "docx":
                try:
                    doc = Document(self.input_file.path)
                    doc.save(pdf_path)
                    if os.path.exists(pdf_path):
                        self.output_file.name = pdf_path
                        self.converted = True
                        self.error_message = None
                    else:
                        raise FileNotFoundError("PDF not created")
                except Exception as e:
                    raise ValueError(f"Error during DOCX to PDF conversion: {e}")

            elif self.source_format == "odt":
                # Conversion ODT en PDF via LibreOffice
                input_path = os.path.abspath(self.input_file.path)
                output_path = os.path.abspath(pdf_path)
                self._convert_odt_to_pdf(input_path, output_path)

            elif self.source_format == "txt":
                # Conversion TXT en PDF
                with open(self.input_file.path, 'r') as txt_file:
                    from fpdf import FPDF 
                    pdf = FPDF()
                    pdf.add_page()
                    pdf.set_font("Arial", size=12)
                    for line in txt_file:
                        pdf.cell(200, 10, txt=line.strip(), ln=True)
                    pdf.output(pdf_path)

            elif self.source_format == "xlsx":
                import pandas as pd
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=12)
                try:
                    data = pd.ExcelFile(self.input_file.path)
                    for sheet_name in data.sheet_names:
                        sheet = data.parse(sheet_name)
                        pdf.cell(200, 10, txt=f"Sheet: {sheet_name}", ln=True)
                        for row in sheet.itertuples(index=False):
                            row_text = " | ".join(map(str, row))
                            pdf.cell(200, 10, txt=row_text, ln=True)
                    pdf.output(pdf_path)
                except Exception as e:
                    self.output_file.name = pdf_path
                    self.converted = True
                    self.error_message = None
                    self.save()
                    raise ValueError(f"Error during XLSX to PDF conversion: {e}")
            
            elif self.source_format == "pptx":
                # Conversion PPTX en PDF via LibreOffice
                input_path = os.path.abspath(self.input_file.path)
                output_path = os.path.abspath(pdf_path)
                self._convert_pptx_to_pdf(input_path, output_path)

            if os.path.exists(pdf_path):
                self.output_file.name = pdf_path
                self.converted = True
                self.error_message = None
                self.save()
            else:
                raise FileNotFoundError(f"PDF file not created at {pdf_path}")

        except Exception as e:
            self.converted = False
            self.error_message = f"Error during {self.source_format} to PDF conversion: {e}"
            self.save()

    def _convert_odt_to_pdf(self, input_path, output_path):
        """Convertir un fichier ODT en PDF via LibreOffice sans afficher les messages."""
        try:
            # Commande pour convertir ODT en PDF avec LibreOffice en mode headless
            command = [
                "soffice",
                "--headless",
                "--invisible",
                "--convert-to", "pdf",
                "--outdir", os.path.dirname(output_path),
                input_path,
            ]

            # Exécute la commande tout en supprimant la sortie standard et les erreurs
            with open(os.devnull, 'w') as devnull:
                subprocess.run(command, stdout=devnull, stderr=devnull, check=True)
            devnull.close()
            # Vérifie si le fichier de sortie a été créé
            if not os.path.exists(output_path):
                raise FileNotFoundError(f"Conversion failed: {output_path} not created")

            self.output_file.name = output_path
            self.converted = True
            self.save()

        except Exception as e:
            self.converted = False
            self.error_message = f"Error during ODT to PDF conversion: {e}"
            self.save()
            
    def _convert_pptx_to_pdf(self, input_path, output_path):
        """Convertir un fichier PPTX en PDF via LibreOffice."""
        try:
            command = [
                "soffice",
                "--headless",
                "--invisible",
                "--convert-to", "pdf",
                "--outdir", os.path.dirname(output_path),
                input_path,
            ]

            # Exécuter la commande en mode silencieux
            with open(os.devnull, 'w') as devnull:
                subprocess.run(command, stdout=devnull, stderr=devnull, check=True)

            # Vérifie si le fichier de sortie a été créé
            if not os.path.exists(output_path):
                raise FileNotFoundError(f"PDF file not created from PPTX: {output_path}")

            self.output_file.name = output_path
            self.converted = True
            self.error_message = None
            self.save()

        except Exception as e:
            self.converted = False
            self.error_message = f"Error during PPTX to PDF conversion: {e}"
            self.save()
            
    def _convert_documents_to_document(self):
        try:
            file_name = os.path.basename(self.input_file.name)
            file_root, _ = os.path.splitext(file_name)
            output_path = os.path.join(settings.MEDIA_ROOT, f"converted/{file_root}.{self.target_format}")

            # Conversion via LibreOffice
            input_path = os.path.abspath(self.input_file.path)
            command = [
                "soffice",
                "--headless",
                "--invisible",
                "--convert-to", self.target_format,
                "--outdir", os.path.dirname(output_path),
                input_path,
            ]

            # Exécute la commande tout en supprimant la sortie standard et les erreurs
            with open(os.devnull, 'w') as devnull:
                subprocess.run(command, stdout=devnull, stderr=devnull, check=True)

            # Vérifie si le fichier de sortie a été créé
            if os.path.exists(output_path):
                self.output_file.name = output_path
                self.converted = True
                self.error_message = None
                self.save()
            else:
                raise FileNotFoundError(f"Converted file not created at {output_path}")

        except Exception as e:
            self.converted = False
            self.error_message = f"Error during {self.source_format} to {self.target_format} conversion: {e}"
            self.save()


